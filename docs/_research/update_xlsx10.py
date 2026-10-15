# -*- coding: utf-8 -*-
"""第十次核对修正：主动预警升级（行 104/155/163）。"""
import openpyxl, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
F = 'docs/功能清单验收对照表.xlsx'
wb = openpyxl.load_workbook(F)
ws = wb['验收标注']

FIXES = {
    104: ('部分实现', '2026-10-15升级：待办由「悬停气泡」扩展为「主动弹窗」——工作台加载后对'
          '本轮新出现的待办弹系统级 ElNotification（patientId+code 去重，处理掉再出现才重新弹，'
          '点击跳「今日待办」视图）；待办规则扩至 5 类（SOFA/APACHE/抗感染复评/ARDS 机械通气/GCS 降幅）。'
          '短信/院内通讯软件通道需外部网关，未接'),
    155: ('已实现', '2026-10-15落地：新增待办规则 ARDS_VENTILATED（patient_info.is_ards=1 且机械通气中，'
          '批量查询不 N+1、不受 24h 前置约束），待办徽章/悬停/主动弹窗三通道提醒，'
          '一键跳转 ARDS 监测页（氧合指数分级/肺保护通气/俯卧位记录）'),
    163: ('已实现', '2026-10-15落地：新增待办规则 GCS_DROP_2（今日 Z_ICU_GCS 正式文书中'
          '首查→最新数值总分下降 ≥2 分；插管 VT 记录口径不可比自动跳过，当日少于 2 条可评全记录不报警），'
          '待办/主动弹窗提醒并一键跳转神经监测页'),
}
for r, (status, ev) in FIXES.items():
    ws.cell(row=r, column=5).value = status
    ws.cell(row=r, column=6).value = ev
    ws.cell(row=r, column=7).value = ''
print('修正行:', sorted(FIXES))
wb.save(F)
print('saved', F)
