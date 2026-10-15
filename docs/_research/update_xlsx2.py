# -*- coding: utf-8 -*-
"""连库核实后二次更新对照表：医护床位比等行改为已实现，追加连库验证说明。"""
import openpyxl, re, io, sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
F = 'docs/功能清单验收对照表.xlsx'
wb = openpyxl.load_workbook(F)

UPDATES = [
    (r'ICU医师总数|医师床位比', '已实现', '2026-10-15连库激活：config_staff ICU相关在职医师58人（档案科室+岗位），排班表未启用故无班次过滤；医师床位比规则302÷304自动出数（2024版ICU-02口径）'),
    (r'ICU护士总数|护士床位比', '已实现', '2026-10-15连库激活：config_staff ICU相关在职护士89人（护士/护师岗）；护士床位比规则303÷304自动出数（2024版ICU-03口径）'),
    (r'医护在岗数|在岗人数', '已实现', '2026-10-15连库激活：按人员档案科室+离职日期过滤；排班表（icu_nurse/doctor_working_schedule）实库0行未启用，无法按班次过滤，口径备注已注明'),
]

ws = wb['验收标注']
changed = 0
for row in ws.iter_rows(min_row=2):
    point = str(row[3].value or '') if len(row) > 3 else ''
    if not point:
        continue
    for pat, status, ev in UPDATES:
        if re.search(pat, point):
            row[4].value = status
            row[5].value = ev
            row[6].value = ''
            changed += 1
            break
print('验收标注 二次更新行数:', changed)

ws3 = wb['汇总统计']
r = ws3.max_row + 1
ws3.cell(r, 1, '2026-10-15 连库核实（172.88.1.184:34567）：确诊表已建、29条指标定义已刷DB、CAUTI规则已插、五个事实层SQL实测通过；医护床位比改自动取数；VAS/FOUR源文书、呼吸机参数、ECMO数据、用药目的字段实库均无（结论见 docs/26 §4）。')

wb.save(F)
print('saved', F)
