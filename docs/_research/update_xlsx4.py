# -*- coding: utf-8 -*-
"""第四次核对修正：本轮交付（标化病死指数/非计划转入/自定义导管类型/危重等级）。"""
import openpyxl, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
F = 'docs/功能清单验收对照表.xlsx'
wb = openpyxl.load_workbook(F)
ws = wb['验收标注']

# 行号 -> (状态, 证据)
FIXES = {
    27: ('已实现', '标化病死指数已自动计算：规则层「标化病死指数=quality_30÷quality_31」（rule 1483050943291009），'
         'quality_30 ICU实际病死率、quality_31 APACHEⅡ预计病死率均 IMPL（引擎取 mortality_rate 求和），重算后自动出数，无需人工录入'),
    30: ('已实现', '2026-10-15落地：quality_20「术后非计划转入例数」接源字段 patient_info.is_plan（0=非计划转入，业务方确认），'
         'where=post_op=1 AND is_plan=0 AND admitted_in_period=1，expr_version+1 已同步入库（IMPL）；'
         '非计划转入率=quality_20÷quality_21 规则自动出数（2024版ICU-14口径）'),
    77: ('已实现', '2026-10-15落地：新增 config_tube_group 配置表（sql/53）+ /api/tube/groups 接口 + 导管统计页「导管类型管理」弹窗，'
         '任意 tube_code 可归入自定义分组并标注危重等级，归类优先级=配置>系统默认白名单>其他，保存后统计即时生效'),
    80: ('部分实现', '2026-10-15落地：月度明细与导管类型管理均带危重等级标注（高/中/低，配置优先，默认气道/CVC=高、导尿管=中、其余=低）；'
         '「按危重等级分模块置顶展示」未做，如需可基于 riskLevel 字段扩展'),
}
for r, (status, ev) in FIXES.items():
    ws.cell(row=r, column=5).value = status
    ws.cell(row=r, column=6).value = ev
    ws.cell(row=r, column=7).value = ''
print('修正行:', sorted(FIXES))
wb.save(F)
print('saved', F)
