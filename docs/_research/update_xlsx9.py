# -*- coding: utf-8 -*-
"""第九次核对修正：统一自定义时间区间组件落地（行 61/66/84）。"""
import openpyxl, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
F = 'docs/功能清单验收对照表.xlsx'
wb = openpyxl.load_workbook(F)
ws = wb['验收标注']

FIXES = {
    61: ('已实现', '质控看板已有「自定义」周期档位（periodType=CUSTOM 起止日期，后端 PeriodRange.CUSTOM）；'
         '指标详情抽屉含患者明细四档视角（全部/分子/分母/未达标）——48h 重返（quality_22）在自定义周期下'
         '点分子即得重返病例清单（/api/quality/metric/patients），可结合患者字段分析特征'),
    66: ('部分实现', '2026-10-15统一组件后：Apache2 总览页接入 TimeRangeBar（本月/近30天/近90天/本年快捷档 + '
         '任意起止日期），分布/平均预期死亡率可按任意区间分析；页内未单独展示实际死亡数与标化死亡指数'
         '（SMR 已由质控看板 quality_30/31 规则自动出数）'),
    84: ('已实现', '2026-10-15落地：导管统计按日视图支持自定义起止区间（TimeRangeBar，本月/近7天/近30天/本年快捷档，'
         '后端 /api/tube/daily startDate/endDate，上限92天），日粒度统计三管新置/拔管/带出/非计划/24h再插管/在管数；'
         '留置天数与感染口径在月度视图（avgDays/totalDays/感染说明）'),
}
for r, (status, ev) in FIXES.items():
    ws.cell(row=r, column=5).value = status
    ws.cell(row=r, column=6).value = ev
    ws.cell(row=r, column=7).value = ''
print('修正行:', sorted(FIXES))
wb.save(F)
print('saved', F)
