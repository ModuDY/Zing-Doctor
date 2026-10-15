# -*- coding: utf-8 -*-
"""按 2026-10-15 缺口落地结果更新《功能清单验收对照表.xlsx》。
只改：验收标注.验收状态/现有实现依据/缺口说明，质控指标数据层.状态/取数实现，
并在汇总统计里追加本轮更新说明。原文件备份为 .bak。"""
import openpyxl, re, io, sys, shutil

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
F = 'docs/功能清单验收对照表.xlsx'
shutil.copyfile(F, F + '.bak')
wb = openpyxl.load_workbook(F)

IMPL_EVIDENCE = [
    # (功能点匹配正则, 状态, 依据文本)
    (r'非计划气管插管拔管', '已实现', '2026-10-15落地：fact_device接patient_tube（is_not_plan_extubation），2024版ICU-12口径；sql/50_quality_gap_impl.sql'),
    (r'48h内再插管|再插管率|48h再插管', '已实现', '2026-10-15落地：fact_device（is_reintubation24，字段名为24实际口径48h），ICU-13口径分子不含非计划拔管后再插管'),
    (r'气管插管拔管数', '已实现', '2026-10-15落地：fact_device（本周期内拔管例数），ICU-12/13共同分母'),
    (r'呼吸机相关肺炎|VAP', '已实现', '2026-10-15落地：新建导管相关感染确诊录入（VAP/CRBSI/CAUTI三型，patient_doc_device_infection），fact_device按导管类型关联，窗口=置入~拔除+48h；需先录入确诊'),
    (r'血流感染|CRBSI', '已实现', '2026-10-15落地：同上，医生确诊录入+fact_device关联（CVC类型），发病率规则ICU侧已有自动出数'),
    (r'导尿管相关泌尿系|CAUTI', '已实现', '2026-10-15落地：同上；CAUTI发病率规则已补（LOCAL_CAUTI_RATE=425÷426×1000‰，2015版十五口径）'),
    (r'血管内导管留置天数|导管留置天数', '已实现', '2026-10-15落地：fact_device.device_days按「本周期∩留置区间」切分求和，未拔截止转出带管/周期末（ICU-17分母）'),
    (r'有创机械通气天数', '已实现', '2026-10-15落地：以气道导管留置天数代理（呼吸机使用记录接入后替换事实层一列）；ICU-16分母'),
    (r'导尿管留置总天数', '已实现', '2026-10-15落地：fact_device.device_days（导尿管类型），2015版十五分母'),
    (r'进行深静脉血栓预防|深静脉血栓预防', '已实现', '2026-10-15落地：fact_dvt接医嘱侧（药物patient_advice_execute_drug/机械与滤器patient_advice_original），关键词按2024版ICU-07官方口径；预防率规则14÷403自动出数'),
    (r'仅进行深静脉血栓药物预防|药物预防的患者总数', '已实现', '2026-10-15落地：fact_dvt.dvt_only_drug/dvt_drug_status（肝素/低分子肝素/利伐沙班/达比加群酯/阿哌沙班/磺达肝癸）'),
    (r'仅进行深静脉血栓机械预防|机械预防的患者总数', '已实现', '2026-10-15落地：fact_dvt.dvt_only_mech/dvt_mech_status（肢体加压泵/气压治疗/梯度压力弹力袜等，ICU-07官方口径）'),
    (r'下腔静脉滤器', '已实现', '2026-10-15落地：fact_dvt.dvt_only_filter/dvt_filter_status（patient_advice_original关键词）'),
    (r'进行深静脉血栓预防并诊断|预防诊断DVT并死亡|预防并死亡', '已实现', '2026-10-15落地：fact_dvt组合列（dvt_prevent_status×dvt_diag/died_in_period）'),
    (r'同期ICU实际开放总床日数|开放总床日', '已实现', '2026-10-15落地：fact_bed激活（config_bed，begin_use_time/stop_use_time），bed_days按周期∩开放区间切分；床位使用率规则2÷301自动出数（ICU-01）'),
    (r'实际开放床位数', '已实现', '2026-10-15落地：fact_bed.open_in_period（本周期内开放状态的床位），ICU-02/03分母'),
    (r'ARDS平均机械通气支持天数', '已实现', '2026-10-15落地：fact_ards.mv_days（按周期切分），AVG=SUM/有创通气ARDS患者数；以插管天数代理'),
    (r'ARDS平均俯卧位通气天数', '已实现', '2026-10-15落地：fact_ards.prone_min（已提交俯卧位记录时长），分母限定实施过俯卧位者'),
    (r'ARDS患者的院内死亡|轻度ARDS患者院内死亡|中度ARDS患者院内死亡|重度ARDS患者院内死亡', '已实现', '2026-10-15落地：按ICU转归死亡（out_vest_type=3）出数（国家ICU-11口径以ICU死亡计）；严格院内死亡需病案首页出院结局'),
    (r'VAS', '部分实现', '2026-10-15落地：镇痛镇静页新增VAS区块（解析/展示/明细表，解析口径同NRS 0-10）；源文书编码Z_ICU_SCORE_VAS待确认，接入后出数'),
    (r'FOUR', '部分实现', '2026-10-15落地：神经监测页新增FOUR区块（0-16，插管可评，2024版ICU-18认可GCS或FOUR）；源文书编码Z_ICU_FOUR待确认'),
    (r'年度视图|按年查询', '已实现', '2026-10-15落地：导管统计页新增年度视图（12月×类型趋势图+矩阵表，/api/tube/yearly，口径与月度一致）'),
    (r'导出', '已实现', '2026-10-15落地：导管统计月度/年度Excel导出（/api/tube/export/*，复用XlsxStreamWriter流式写出）'),
    (r'按日', '部分实现', '2026-10-15落地：月度/年度视图与导出已完成；按日使用明细尚未实现'),
    (r'院感确诊录入|VAP.{0,6}CRBSI.{0,6}CAUTI确诊', '已实现', '2026-10-15落地：新页「导管相关感染确诊」（临床分析菜单），三型共用表单，科室权限校验同查房记录；驱动quality_24/26/425'),
]

# ---- Sheet: 验收标注 ----
ws = wb['验收标注']
changed = 0
for row in ws.iter_rows(min_row=2):
    point = str(row[3].value or '') if len(row) > 3 else ''
    if not point:
        continue
    for pat, status, ev in IMPL_EVIDENCE:
        if re.search(pat, point):
            row[4].value = status
            row[5].value = ev
            row[6].value = ''
            changed += 1
            break
print('验收标注 更新行数:', changed)

# ---- Sheet: 质控指标数据层 ----
IMPL_CODES = {'quality_17','quality_18','quality_19','quality_24','quality_25','quality_26','quality_27',
              'quality_425','quality_426','quality_14','quality_404','quality_405','quality_406','quality_407',
              'quality_408','quality_409','quality_411','quality_412','quality_413','quality_301','quality_304',
              'quality_477','quality_486','quality_487','quality_488','quality_489','quality_490'}
ws2 = wb['质控指标数据层']
c2 = 0
for row in ws2.iter_rows(min_row=2):
    code = str(row[1].value or '')
    if code in IMPL_CODES:
        row[4].value = 'IMPL'
        row[7].value = '2026-10-15落地（见 docs/26）'
        c2 += 1
print('质控指标数据层 更新行数:', c2)

# ---- Sheet: 汇总统计 追加说明 ----
ws3 = wb['汇总统计']
r = ws3.max_row + 2
ws3.cell(r, 1, '2026-10-15 缺口落地：27条指标IMPL（插管3/导管感染与天数6/DVT预防10/床位2/ARDS6）+ 导管感染确诊录入页 + 导管统计年度视图与导出 + VAS/FOUR区块。口径依据《重症医学专业医疗质量控制指标》2015/2024版（docs/26）。抗菌药治疗性用药、医护床位比、保护性肺通气、ECMO天数、设备管理、彩超教学等仍待源数据/口径确认（清单见 docs/26 §4）。')

wb.save(F)
print('saved', F)
