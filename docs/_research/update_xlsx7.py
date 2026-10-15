# -*- coding: utf-8 -*-
"""第七次核对修正：按日导管统计、感染控制分析页。"""
import openpyxl, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
F = 'docs/功能清单验收对照表.xlsx'
wb = openpyxl.load_workbook(F)
ws = wb['验收标注']

FIXES = {
    79: ('已实现', '2026-10-15落地：导管统计页新增「按日」视图（/api/tube/daily + TubeStatsMapper.dailyRows），'
         '按「日 × 导管分组」展示每天新置管/拔管/带出/非计划拔管/24h再插管，未拔导管算至月末；'
         '月度/年度/按日三视图可切'),
    122: ('已实现', '2026-10-15落地：院感分析页 /page/infection-watch（sql/54 注册），'
          '体温（38℃标线）+ WBC/PCT/CRP/IL-6 感染指标联合趋势（/api/lab-trend?type=infection 新面板，'
          '按项目名匹配、排除体液标本白细胞），附检验明细表'),
}
for r, (status, ev) in FIXES.items():
    ws.cell(row=r, column=5).value = status
    ws.cell(row=r, column=6).value = ev
    ws.cell(row=r, column=7).value = ''
print('修正行:', sorted(FIXES))
wb.save(F)
print('saved', F)
