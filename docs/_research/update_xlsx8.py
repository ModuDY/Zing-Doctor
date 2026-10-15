# -*- coding: utf-8 -*-
"""第八次核对修正：感染性休克 bundle 近一年完成率（驾驶舱已实现）。"""
import openpyxl, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
F = 'docs/功能清单验收对照表.xlsx'
wb = openpyxl.load_workbook(F)
ws = wb['验收标注']
ws.cell(row=53, column=5).value = '已实现'
ws.cell(row=53, column=6).value = ('驾驶舱「SEPSIS BUNDLE」面板：近 12 个月感染性休克 bundle 完成率折线'
    '（quality_11÷quality_10，bundleTrend 逐月 分子/分母/完成率，缺月断开不按 0）；'
    'fact_sepsis_bundle + 脓毒症域 26 项 IMPL，数据由每日批算落库')
ws.cell(row=53, column=7).value = ''
print('修正行: [53]')
wb.save(F)
print('saved', F)
