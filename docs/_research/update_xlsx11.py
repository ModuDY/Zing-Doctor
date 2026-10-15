# -*- coding: utf-8 -*-
"""第十一次核对修正：VAS/FOUR 不适用、柏林定义已覆盖、设备状态页落地（行 51/144/150/156/161）。"""
import openpyxl, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
F = 'docs/功能清单验收对照表.xlsx'
wb = openpyxl.load_workbook(F)
ws = wb['验收标注']

FIXES = {
    51: ('已实现', '2026-10-15落地：设备状态页 /page/device-status（sql/57 注册 + GET /api/department-cockpit/device-status），'
         '只读重症系统 config_device：设备编码/类型（监护仪/呼吸机/CRRT/ECMO）/科室/当前绑定患者（在科四类设备编码精确匹配）/'
         '采集连接状态（正在接收/已连接/初始化/停止/异常/未接入）/最近更新，顶部六卡汇总（配置/使用中/未绑定/接收/异常/未接入），'
         '类型筛选；科室权限与驾驶舱同套路；口径与驾驶舱「设备接入态势」同源。源系统无维护台账，页面明确不虚构维修/库存语义'),
    144: ('不适用（源系统未启用）', '2026-10-15业务方确认：重症系统暂未使用 VAS 镇痛评分文书，不出数也无需再接。'
          '页面侧 VAS 区块保留（Z_ICU_SCORE_VAS 编码约定），源系统启用后仅需确认编码即可点亮；NRS 已实现'),
    150: ('已实现', '2026-10-15业务方确认：柏林定义的质控语义已由氧合指数（OI）自动分级覆盖——ARDS 监测页按'
          '氧合指数自动分级轻度/中度/重度并特别标识，ARDS 全院筛查页按诊断清单化；视为已实现'),
    156: ('不适用（源系统未启用）', '2026-10-15业务方确认：重症系统暂未使用 VAS 文书。NRS 已实现（镇痛镇静页），'
          'VAS 区块保留待源系统启用；多项标准化评估工具的「规范性/可比性」以 NRS + CPOT + CAM-ICU 承载'),
    161: ('不适用（源系统未启用）', '2026-10-15业务方确认：重症系统暂未使用 FOUR 评分文书，不出数也无需再接。'
          '页面侧 FOUR 区块保留（Z_ICU_FOUR 编码约定），源系统启用后仅需确认编码即可点亮；GCS + 瞳孔 + ICP + 神志已实现'),
}
for r, (status, ev) in FIXES.items():
    ws.cell(row=r, column=5).value = status
    ws.cell(row=r, column=6).value = ev
    ws.cell(row=r, column=7).value = ''
print('修正行:', sorted(FIXES))
wb.save(F)
print('saved', F)
