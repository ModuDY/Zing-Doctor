# -*- coding: utf-8 -*-
"""第六次核对修正：血流动力学/心功能、ADL、床旁超声教学。"""
import openpyxl, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
F = 'docs/功能清单验收对照表.xlsx'
wb = openpyxl.load_workbook(F)
ws = wb['验收标注']

FIXES = {
    67: ('已实现', '2026-10-15落地：GCS / RASS 已有（神经监测/镇静镇静页），'
         'ADL（Barthel 指数）新增手工录入模块（patient_doc_adl 表 + /api/adl/* + /page/hemodynamics 第二区块），'
         '十项 0/5/10/15 逐项录入、总分服务端求和、依赖四级分色 + 总分趋势图'),
    114: ('已实现', '2026-10-15落地：血流动力学手工录入模块（patient_doc_hemodynamics 表 + /api/hemodynamics/*），'
          '录入 SBP/DBP/MAP/CVP/CO，MAP 缺省按 DBP+(SBP-DBP)/3 补算、SVR=(MAP-CVP)/CO×80 服务端自动计算，'
          '多 y 轴趋势图（压力/CO/SVR+EF）+ 明细表 + 编辑删除'),
    116: ('已实现', '2026-10-15落地：心功能（EF 射血分数）随血流动力学模块手工录入并趋势展示'
          '（源系统无 CO/CVP/EF 结构化数据，docs/25 §3 连库核查结论；页面标注「EF 仅展示，不做心衰结论」），'
          '测量方式区分 PiCCO / Swan-Ganz / 床旁超声'),
    143: ('已实现', '2026-10-15落地：床旁超声教学静态页 /page/echo-teaching（sql/56 注册），'
          '八部位（肺部/眼球/下腔静脉/腹部/膈肌/肾脏/心脏/血管）各含检查目的/操作要点/正常异常判读/注意事项，'
          '支持关键词过滤；标注「仅供教学参考，正式诊断以超声科报告为准」'),
}
for r, (status, ev) in FIXES.items():
    ws.cell(row=r, column=5).value = status
    ws.cell(row=r, column=6).value = ev
    ws.cell(row=r, column=7).value = ''
print('修正行:', sorted(FIXES))
wb.save(F)
print('saved', F)
